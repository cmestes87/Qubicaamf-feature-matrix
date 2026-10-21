"""
Rebuilds source/data.json from the two source spreadsheets.

Run this after replacing either spreadsheet in source/spreadsheets/:
    cd source
    python3 extract.py

Then rebuild the deployable site with:
    python3 build.py
"""
import openpyxl, json, os

BASE = os.path.dirname(os.path.abspath(__file__))
NEOVERSE_XLSX = os.path.join(BASE, "spreadsheets", "Neoverse_for_Claude.xlsx")
BESNV_XLSX = os.path.join(BASE, "spreadsheets", "BES_NV_Environments_and_Features_for_Claude.xlsx")
OUT = os.path.join(BASE, "data.json")


def clean(v):
    if v is None:
        return None
    if isinstance(v, str):
        return v.strip()
    return v


# ---- Neoverse ----
# Sheet1 (2), columns: Name | Neoverse Interactive Edition | Neoverse (Stand Alone) | Requires Subscription | Requires BES NV | Requires SuperTouch | Content | Feature | New
# Category header rows have a name in col A and every other column empty.
wb = openpyxl.load_workbook(NEOVERSE_XLSX, data_only=True)
ws = wb.worksheets[0]
rows = list(ws.iter_rows(values_only=True))

neoverse_categories = []
cur = None
for r in rows:
    name, interactive, standalone, sub, reqbes, supertouch, content, feature, new = [clean(x) for x in r]
    if name is None and interactive is None:
        continue  # blank separator row
    if (interactive is None and standalone is None and sub is None and reqbes is None
            and supertouch is None and content is None and feature is None and new is None):
        cur = {"name": name, "items": []}
        neoverse_categories.append(cur)
        continue
    if cur is None:
        continue
    cur["items"].append({
        "name": name,
        "interactive": interactive == "Included",
        "standalone": standalone == "Included",
        "sub": sub == "Yes",
        "superTouch": supertouch == "Yes",
        "neofi": reqbes == "Yes",
        "isFeature": bool(feature),
        "isNew": bool(new),
    })

# ---- BES NV ----
# Sheet2. Columns are located by HEADER NAME (not position), so inserting a new column
# anywhere won't silently shift the data. Expected headers:
#   Environments | BES NV | Requires Subscription | BES NV Value | Requires Subscription |
#   Added After Launch | Neo-Fi | Requires SuperTouch | Environment | Feature | New
# "Requires Subscription" appears twice: the first belongs to BES NV, the second to BES NV Value.
# Category header rows (e.g. "HyperBowling", "Classic Games") have a name in col A and every other column empty.
# Rows before the first named category header belong to "Environments".
wb2 = openpyxl.load_workbook(BESNV_XLSX, data_only=True)
ws2 = wb2["Sheet2"]
rows2 = list(ws2.iter_rows(values_only=True))

header_row = next(r for r in rows2 if r and r[0] == "Environments" and r[1] == "BES NV")
cols = {}
for idx, h in enumerate(header_row):
    if h is None:
        continue
    cols.setdefault(str(h).strip(), []).append(idx)

def col(name, nth=0):
    try:
        return cols[name][nth]
    except (KeyError, IndexError):
        return None

C_BESNV, C_BESNV_SUB = col("BES NV"), col("Requires Subscription", 0)
C_VALUE, C_VALUE_SUB = col("BES NV Value"), col("Requires Subscription", 1)
C_NEOFI, C_TOUCH = col("Neo-Fi"), col("Requires SuperTouch")
C_FEATURE, C_NEW = col("Feature"), col("New")
C_LAUNCH = col("Added After Launch")  # optional; absent in older sheets
required = {"BES NV": C_BESNV, "BES NV Value": C_VALUE, "Neo-Fi": C_NEOFI,
            "Requires SuperTouch": C_TOUCH, "Feature": C_FEATURE, "New": C_NEW}
missing = [k for k, v in required.items() if v is None]
if missing or C_BESNV_SUB is None or C_VALUE_SUB is None:
    raise SystemExit(f"BES NV sheet is missing expected columns: {missing or 'Requires Subscription x2'}")

def cell(r, i):
    return clean(r[i]) if i is not None and i < len(r) else None

besnv_categories = []
cur = None
for r in rows2:
    name = cell(r, 0)
    if name is None or r is header_row:
        continue
    flag_values = [cell(r, i) for i in range(1, len(r))]
    if all(v is None for v in flag_values):
        cur = {"name": name, "items": []}
        besnv_categories.append(cur)
        continue
    if cur is None:
        cur = {"name": "Environments", "items": []}
        besnv_categories.append(cur)
    cur["items"].append({
        "name": name,
        "besnv": bool(cell(r, C_BESNV)),
        "besnvSub": bool(cell(r, C_BESNV_SUB)),
        "value": bool(cell(r, C_VALUE)),
        "valueSub": bool(cell(r, C_VALUE_SUB)),
        "neofi": bool(cell(r, C_NEOFI)),
        "superTouch": bool(cell(r, C_TOUCH)),
        "isFeature": bool(cell(r, C_FEATURE)),
        "isNew": bool(cell(r, C_NEW)),
        "addedAfterLaunch": bool(cell(r, C_LAUNCH)),
    })

out = {"neoverse": neoverse_categories, "besnv": besnv_categories}
with open(OUT, "w", encoding="utf-8") as f:
    json.dump(out, f, indent=2, ensure_ascii=False)

print("BES NV categories:", [(c["name"], len(c["items"])) for c in besnv_categories])
print("Neoverse categories:", [(c["name"], len(c["items"])) for c in neoverse_categories])
print(f"Wrote {OUT}")
